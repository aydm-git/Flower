from flask import Flask, render_template, request
import openpyxl
import os
from datetime import datetime

app = Flask(__name__)

EXCEL_FILE = "orders.xlsx"

products = [
    {'id': 1, 'name': 'Букет роз', 'price': 2500, 'img': '/static/images/roses.jpeg'},
    {'id': 2, 'name': 'Тюльпаны', 'price': 1800, 'img': '/static/images/tulps.jpg'},
    {'id': 3, 'name': 'Микс-букет', 'price': 3200, 'img': '/static/images/mix.jpeg'}, 
    {'id': 4, 'name': 'Цветы в коробке', 'price': 4000, 'img': '/static/images/box.jpg'},
]

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Заказы"
        ws.append(["ID", "Дата", "Имя", "Телефон", "Букет"])
        wb.save(EXCEL_FILE)

def save_to_excel(name, phone, flower):
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        next_id = ws.max_row 
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.append([next_id, now, name, phone, flower])
        wb.save(EXCEL_FILE)
        return True
    except Exception as e:
        print(f"[EXCEL ERROR] {e}")
        return False

init_excel()

@app.route("/", methods=["GET", "POST"])
def home():
    message = ""
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        phone = request.form.get("phone", "").strip()
        flower = request.form.get("flower", "").strip()

        if not name or not phone:
            message = "Заполните имя и телефон"
        else:
            if save_to_excel(name, phone, flower):
                message = f"Спасибо, {name}! Заявка на «{flower}» принята."
            else:
                message = "Произошла ошибка. Попробуйте позже."

    return render_template("index.html", products=products, message=message)

if __name__ == "__main__":
    app.run(debug=True)